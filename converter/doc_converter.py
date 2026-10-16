import os
import subprocess
import json
import csv
import shutil
import tempfile
import markdown
import docx
import openpyxl

def convert_document(input_path, output_path, src_format, target_format, options=None):
    src = src_format.lower()
    tgt = target_format.lower()

    if src in ("docx", "odt", "rtf"):
        if tgt == "pdf":
            return libreoffice_convert(input_path, output_path, "pdf")
        elif tgt == "html":
            return libreoffice_convert(input_path, output_path, "html")
        elif tgt == "txt":
            if src == "docx":
                return docx_to_txt(input_path, output_path)
            return libreoffice_convert(input_path, output_path, "txt:Text")
        elif tgt in ("docx", "odt"):
            return libreoffice_convert(input_path, output_path, tgt)

    if src == "md":
        with open(input_path, "r", encoding="utf-8", errors="replace") as f:
            md_content = f.read()
        html_content = markdown.markdown(md_content, extensions=['tables', 'fenced_code', 'toc'])

        if tgt == "html":
            full_html = f"<!DOCTYPE html><html><head><meta charset='utf-8'><title>Document</title><style>body{{font-family:sans-serif;max-width:800px;margin:2rem auto;line-height:1.6;padding:0 1rem;}}table{{border-collapse:collapse;width:100%;}}th,td{{border:1px solid #ddd;padding:8px;}}pre{{background:#f4f4f4;padding:12px;border-radius:4px;}}</style></head><body>{html_content}</body></html>"
            with open(output_path, "w", encoding="utf-8") as f:
                f.write(full_html)
            return output_path
        elif tgt == "txt":
            with open(output_path, "w", encoding="utf-8") as f:
                f.write(md_content)
            return output_path
        elif tgt in ("pdf", "docx"):
            temp_html = tempfile.mktemp(suffix=".html")
            full_html = f"<!DOCTYPE html><html><head><meta charset='utf-8'></head><body>{html_content}</body></html>"
            with open(temp_html, "w", encoding="utf-8") as f:
                f.write(full_html)
            try:
                res = libreoffice_convert(temp_html, output_path, tgt)
                return res
            finally:
                if os.path.exists(temp_html):
                    os.remove(temp_html)

    if src == "html":
        if tgt == "pdf":
            return libreoffice_convert(input_path, output_path, "pdf")
        elif tgt == "txt":
            with open(input_path, "r", encoding="utf-8", errors="replace") as f:
                html_text = f.read()
            import re
            clean_text = re.sub(r'<style.*?</style>', '', html_text, flags=re.DOTALL)
            clean_text = re.sub(r'<script.*?</script>', '', clean_text, flags=re.DOTALL)
            clean_text = re.sub(r'<[^>]+>', '', clean_text)
            clean_text = re.sub(r'\n\s*\n', '\n\n', clean_text).strip()
            with open(output_path, "w", encoding="utf-8") as f:
                f.write(clean_text)
            return output_path

    if src == "txt":
        with open(input_path, "r", encoding="utf-8", errors="replace") as f:
            content = f.read()
        if tgt == "html":
            import html
            safe_text = html.escape(content).replace("\n", "<br/>")
            with open(output_path, "w", encoding="utf-8") as f:
                f.write(f"<!DOCTYPE html><html><head><meta charset='utf-8'></head><body><pre>{safe_text}</pre></body></html>")
            return output_path
        elif tgt == "md":
            with open(output_path, "w", encoding="utf-8") as f:
                f.write(content)
            return output_path
        elif tgt == "pdf":
            temp_txt = tempfile.mktemp(suffix=".txt")
            with open(temp_txt, "w", encoding="utf-8") as f:
                f.write(content)
            try:
                return libreoffice_convert(temp_txt, output_path, "pdf")
            finally:
                if os.path.exists(temp_txt):
                    os.remove(temp_txt)

    if src == "csv" and tgt == "xlsx":
        wb = openpyxl.Workbook()
        ws = wb.active
        with open(input_path, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.reader(f)
            for row in reader:
                ws.append(row)
        wb.save(output_path)
        return output_path

    if src == "xlsx" and tgt == "csv":
        wb = openpyxl.load_workbook(input_path, data_only=True)
        ws = wb.active
        with open(output_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                writer.writerow(["" if c is None else str(c) for c in row])
        return output_path

    if src == "csv" and tgt == "json":
        with open(input_path, "r", encoding="utf-8", errors="replace") as f:
            reader = csv.DictReader(f)
            rows = list(reader)
        with open(output_path, "w", encoding="utf-8") as f:
            json.dump(rows, f, indent=2, ensure_ascii=False)
        return output_path

    if src == "json" and tgt == "csv":
        with open(input_path, "r", encoding="utf-8") as f:
            data = json.load(f)

        if isinstance(data, dict):
            for v in data.values():
                if isinstance(v, list) and len(v) > 0 and isinstance(v[0], dict):
                    data = v
                    break
            else:
                data = [{"Key": k, "Value": str(v)} for k, v in data.items()]

        if not isinstance(data, list) or len(data) == 0:
            raise ValueError("JSON data must contain an array of objects to convert to CSV.")

        fieldnames = []
        for item in data:
            if isinstance(item, dict):
                for k in item.keys():
                    if k not in fieldnames:
                        fieldnames.append(k)

        with open(output_path, "w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=fieldnames)
            writer.writeheader()
            for item in data:
                if isinstance(item, dict):
                    writer.writerow(item)
        return output_path

    if src == "json" and tgt == "txt":
        with open(input_path, "r", encoding="utf-8") as f:
            data = json.load(f)
        with open(output_path, "w", encoding="utf-8") as f:
            f.write(json.dumps(data, indent=2, ensure_ascii=False))
        return output_path

    raise NotImplementedError(f"Conversion from {src} to {tgt} is not supported.")

def libreoffice_convert(input_path, output_path, out_format):
    out_dir = os.path.dirname(output_path)
    cmd = [
        "soffice",
        "--headless",
        "--convert-to", out_format,
        "--outdir", out_dir,
        input_path
    ]
    res = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    if res.returncode != 0:
        raise RuntimeError(f"LibreOffice conversion failed: {res.stderr}")

    base_name = os.path.splitext(os.path.basename(input_path))[0]
    expected_ext = out_format.split(":")[0]
    produced_file = os.path.join(out_dir, f"{base_name}.{expected_ext}")

    if os.path.exists(produced_file) and produced_file != output_path:
        shutil.move(produced_file, output_path)

    if not os.path.exists(output_path) or os.path.getsize(output_path) == 0:
        raise RuntimeError(f"LibreOffice conversion did not produce valid output at {output_path}")

    return output_path

def docx_to_txt(input_path, output_path):
    doc = docx.Document(input_path)
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            paragraphs.append("\t".join(cell.text.strip() for cell in row.cells))

    with open(output_path, "w", encoding="utf-8") as f:
        f.write("\n\n".join(paragraphs))

    return output_path
